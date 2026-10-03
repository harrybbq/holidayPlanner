"""Import scripts/holiday_planner.xlsx into data/seed.json (conforms to SeedFile in src/types.ts).

The work travel clearance lists go to data/countries.json instead. That file and the xlsx are
local only (gitignored), so the committed seed carries no clearance data.

Only literal input cells are read; formula cells (Nights, Shifts, totals) are ignored and the
shift counts are recomputed here for verification. Re-runnable: overwrites data/seed.json and
data/countries.json only.

Usage: python scripts/import_xlsx.py
"""

from __future__ import annotations

import json
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent
XLSX = SCRIPT_DIR / "holiday_planner.xlsx"
OUT = ROOT / "data" / "seed.json"
COUNTRIES_OUT = ROOT / "data" / "countries.json"

IN_PLAN = {"Yes", "Parked"}
TRIP_STATUS = {"Idea", "Planning", "Booked", "Done", "Cancelled"}
LEAVE_STATUS = {"Not requested", "Requested", "Approved", "Declined", "Cancelled"}
CLEARANCE_STATUS = {"Pre-cleared", "Request needed", "Likely refused", "Unknown"}

# name -> (country, lat, lng). Country names match the Countries sheet where one exists.
PLACES: dict[str, tuple[str, float, float]] = {
    "Budapest": ("Hungary", 47.498, 19.040),
    "Rome": ("Italy", 41.903, 12.496),
    "Venice": ("Italy", 45.441, 12.316),
    "Verona": ("Italy", 45.438, 10.992),
    "Como": ("Italy", 45.808, 9.085),
    "Milan": ("Italy", 45.464, 9.190),
    "Prague": ("Czechia", 50.076, 14.438),
    "Split": ("Croatia", 43.508, 16.440),
    "Dubrovnik": ("Croatia", 42.650, 18.094),
    "Kotor": ("Montenegro", 42.425, 18.771),
    "Oslo": ("Norway", 59.914, 10.752),
    "Lofoten": ("Norway", 68.234, 14.568),  # Svolvær
    "New York": ("United States", 40.713, -74.006),
    "Hudson Valley": ("United States", 41.700, -73.921),  # Poughkeepsie
    "New England": ("United States", 43.624, -72.518),  # Woodstock, VT
    "Boston": ("United States", 42.360, -71.058),
    "Seattle": ("United States", 47.606, -122.332),
    "Portland": ("United States", 45.515, -122.679),  # Oregon
    "San Francisco": ("United States", 37.775, -122.419),
    "Yosemite": ("United States", 37.745, -119.593),  # Yosemite Valley
    "Las Vegas": ("United States", 36.170, -115.140),
    "Dublin": ("Ireland", 53.350, -6.260),
    "Munich": ("Germany", 48.137, 11.576),
    "Kaltenberg": ("Germany", 48.176, 11.009),
    "Salzburg": ("Austria", 47.809, 13.055),
    "Innsbruck": ("Austria", 47.269, 11.404),
    "Tokyo": ("Japan", 35.676, 139.650),
    "Kyoto": ("Japan", 35.012, 135.768),
    "Mt Fuji": ("Japan", 35.361, 138.727),
    "Sydney": ("Australia", -33.869, 151.209),
    "Gold Coast": ("Australia", -28.017, 153.400),
}

# Destination token (or whole destination) -> canonical PLACES key, or None to drop the token.
ALIASES: dict[str, str | None] = {
    "Budapest, Hungary": "Budapest",
    "Prague, Czechia": "Prague",
    "Lake Como": "Como",
    "castles": None,
    "Mount Fuji": "Mt Fuji",
}


def text(v) -> str:
    return "" if v is None else str(v).strip()


def iso(v) -> str | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    raise ValueError(f"Expected a date, got {v!r}")


def slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def number(v) -> float | int:
    if v is None or v == "":
        return 0
    n = float(v)
    return int(n) if n.is_integer() else n


def check(value: str, allowed: set[str], what: str, where: str) -> str:
    if value not in allowed:
        raise ValueError(f"{where}: {what} {value!r} not in {sorted(allowed)}")
    return value


def header_row(ws, first: str) -> tuple[int, dict[str, int]]:
    """Find the row whose column A equals `first`; return (row, {header: col})."""
    for r in range(1, ws.max_row + 1):
        if text(ws.cell(r, 1).value) == first:
            cols = {}
            for c in range(1, ws.max_column + 1):
                h = text(ws.cell(r, c).value)
                if h:
                    cols[h] = c
            return r, cols
    raise ValueError(f"{ws.title}: header row starting {first!r} not found")


def col(cols: dict[str, int], prefix: str) -> int:
    # Prefix match so a mangled header like 'Est. cost (�)' still resolves.
    for h, c in cols.items():
        if h.startswith(prefix):
            return c
    raise ValueError(f"Column starting {prefix!r} not found in {list(cols)}")


def read_settings(wb) -> dict:
    ws = wb["Setup"]
    found = {}
    labels = {
        "allowanceShifts": "Annual leave",
        "rotaAnchor": "Rota anchor",
        "monthsUntilLastTrip": "Months until",
    }
    for r in range(1, ws.max_row + 1):
        label = text(ws.cell(r, 2).value)
        for key, prefix in labels.items():
            if key not in found and label.startswith(prefix):
                found[key] = ws.cell(r, 3).value
    missing = set(labels) - set(found)
    if missing:
        raise ValueError(f"Setup: labels not found for {missing}")
    return {
        "allowanceShifts": number(found["allowanceShifts"]),
        "rotaAnchor": iso(found["rotaAnchor"]),
        "monthsUntilLastTrip": number(found["monthsUntilLastTrip"]),
    }


def read_countries(wb) -> list[dict]:
    ws = wb["Countries"]
    hr, cols = header_row(ws, "Country")
    cn, cs, cnote = col(cols, "Country"), col(cols, "Status"), col(cols, "Note")
    out = []
    # The list has blank rows and label rows between sections, so scan to the end and keep
    # only rows that carry a status.
    for r in range(hr + 1, ws.max_row + 1):
        name, status = text(ws.cell(r, cn).value), text(ws.cell(r, cs).value)
        if not name or not status:
            continue
        check(status, CLEARANCE_STATUS, "clearance status", f"Countries!{r}")
        out.append({"name": name, "status": status, "note": text(ws.cell(r, cnote).value)})
    return out


def read_leave(wb) -> dict[str, dict]:
    ws = wb["Leave Tracker"]
    hr, cols = header_row(ws, "Trip")
    ct, cs, creq, cnote = (col(cols, h) for h in ("Trip", "Status", "Requested on", "Notes"))
    out = {}
    for r in range(hr + 1, ws.max_row + 1):
        trip = text(ws.cell(r, ct).value)
        if trip.startswith("Allowance"):
            break
        if not trip:
            continue
        status = text(ws.cell(r, cs).value) or "Not requested"
        check(status, LEAVE_STATUS, "leave status", f"Leave Tracker!{r}")
        out[trip] = {
            "status": status,
            "requestedOn": iso(ws.cell(r, creq).value),
            "notes": text(ws.cell(r, cnote).value),
        }
    return out


def make_stop(name: str, nights: int | None) -> dict:
    country, lat, lng = PLACES[name]
    return {"name": name, "country": country, "lat": lat, "lng": lng, "nights": nights}


def parse_stops(destination: str, notes: str) -> list[dict]:
    if not destination or destination.upper().startswith("TBC"):
        return []
    if destination in ALIASES:
        tokens = [destination]
    else:
        tokens = [t.strip() for t in re.split(r" - | \+ ", destination) if t.strip()]
    names = []
    for t in tokens:
        key = ALIASES.get(t, t)
        if key is None:
            continue
        if key not in PLACES:
            raise ValueError(f"No coordinates for stop {t!r} (destination {destination!r})")
        names.append(key)
    # Nights per stop, only when notes spell them out as "<Stop> <n>" for a stop on this route.
    nights: dict[str, int] = {}
    for word, n in re.findall(r"([A-Za-z][A-Za-z ]*?)\s+(\d+)\b", notes):
        word = word.strip()
        if word in names and word not in nights:
            nights[word] = int(n)
    return [make_stop(n, nights.get(n)) for n in names]


def read_trips(wb, leave: dict[str, dict]) -> list[dict]:
    ws = wb["Trips"]
    hr, cols = header_row(ws, "Trip")
    c = {
        k: col(cols, h)
        for k, h in {
            "name": "Trip", "dest": "Destination", "start": "Start", "end": "End",
            "who": "Who with", "inPlan": "In plan", "cost": "Est. cost", "status": "Status",
            "bookBy": "Book by", "clearance": "Clearance", "notified": "Notified on",
            "notes": "Notes",
        }.items()
    }
    trips = []
    for r in range(hr + 1, ws.max_row + 1):
        name = text(ws.cell(r, c["name"]).value)
        if not name or name.upper() == "IN PLAN":
            break
        v = lambda k: ws.cell(r, c[k]).value  # noqa: E731
        where = f"Trips!{r}"
        destination, notes = text(v("dest")), text(v("notes"))
        trips.append({
            "id": slug(name),
            "name": name,
            "destination": destination,
            "stops": parse_stops(destination, notes),
            "start": iso(v("start")),
            "end": iso(v("end")),
            "whoWith": text(v("who")),
            "inPlan": check(text(v("inPlan")), IN_PLAN, "in plan", where),
            "cost": number(v("cost")),
            "status": check(text(v("status")), TRIP_STATUS, "status", where),
            "bookBy": iso(v("bookBy")),
            # The sheet's free-text Clearance column stays in the xlsx; the app works it out
            # from the stops instead.
            "clearance": "",
            "notifiedOn": iso(v("notified")),
            "notes": notes,
            "leave": leave.get(name, {"status": "Not requested", "requestedOn": None, "notes": ""}),
        })
    ids = [t["id"] for t in trips]
    if len(ids) != len(set(ids)):
        raise ValueError(f"Duplicate trip ids: {ids}")
    unmatched = set(leave) - {t["name"] for t in trips}
    if unmatched:
        print(f"WARNING: leave rows with no matching trip: {sorted(unmatched)}")
    return trips


def shifts(start: str | None, end: str | None, anchor: str) -> int | None:
    """Sheet formula: days d in [start, end] inclusive with ((d - anchor) mod 8) < 4."""
    if not start or not end:
        return None
    s, e, a = (date.fromisoformat(x) for x in (start, end, anchor))
    return sum(1 for i in range((e - s).days + 1) if ((s + timedelta(i) - a).days % 8) < 4)


def verify(seed: dict, countries: list[dict]) -> None:
    anchor = seed["settings"]["rotaAnchor"]
    allowance = seed["settings"]["allowanceShifts"]
    print(f"{'Trip':<16}{'Start':<12}{'End':<12}{'Nights':>7}{'Shifts':>7}  {'In plan':<8}{'Leave':<15}{'Cost':>7}")
    used = cost_plan = cost_parked = 0
    for t in seed["trips"]:
        sh = shifts(t["start"], t["end"], anchor)
        nights = (date.fromisoformat(t["end"]) - date.fromisoformat(t["start"])).days if t["start"] and t["end"] else None
        print(f"{t['name']:<16}{t['start'] or '-':<12}{t['end'] or '-':<12}"
              f"{'-' if nights is None else nights:>7}{'-' if sh is None else sh:>7}  "
              f"{t['inPlan']:<8}{t['leave']['status']:<15}{t['cost']:>7}")
        if t["inPlan"] == "Yes":
            cost_plan += t["cost"]
            if t["leave"]["status"] != "Cancelled":
                used += sh or 0
        else:
            cost_parked += t["cost"]
    print()
    print(f"Allowance: {allowance}  Shifts used (in plan, leave not Cancelled): {used}  Remaining: {allowance - used}")
    print(f"Cost in plan: {cost_plan}  Cost parked: {cost_parked}")
    print(f"Countries: {len(countries)}  Trips: {len(seed['trips'])}")


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    wb = openpyxl.load_workbook(XLSX, data_only=False)
    seed = {
        "settings": read_settings(wb),
        "trips": read_trips(wb, read_leave(wb)),
    }
    countries = read_countries(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(seed, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    COUNTRIES_OUT.write_text(json.dumps(countries, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {OUT}")
    print(f"Wrote {COUNTRIES_OUT}")
    verify(seed, countries)


if __name__ == "__main__":
    main()
